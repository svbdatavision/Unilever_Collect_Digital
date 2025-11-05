import fitz  # PyMuPDF
import pandas as pd
import re
import os
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Border, Side
from tkinter import messagebox

# --- PATRONES Y MAPEO ---
TIPO_MAP = {
    "Fact.Elect. Af. Emi": "Factura por convenio",
    "Fac Elect Ex Emitida": "Factura por convenio",
    "Ndd Afecta Elec. Rec": "Nota de débito",
    "Fac Afecta Elec. Rec": "Factura",
    "Ncr Ex Elect. Rec": "Nota de crédito"
}
PREFIJOS_A_ELIMINAR = ("08-", "07-", "00-", "01-")

# --- NUEVO: DETECCIÓN ROBUSTA DE FILAS Y CAMPOS ---

def extraer_tabla_pdf(archivo_pdf):
    doc = fitz.open(archivo_pdf)
    tabla_documentos = []
    regex_fila = re.compile(
        r"(\d{2}\.\d{2}\.\d{4})\s+([A-Z0-9\-]+)\s+([\w\s\.\-]+?)\s+([\-]?\d{1,3}(?:[\.,]\d{3})*[\.,]\d{2})"
    )
    for page in doc:
        text = page.get_text()
        for match in regex_fila.findall(text):
            fecha, referencia, tipo_doc, monto_texto = match
            for prefijo in PREFIJOS_A_ELIMINAR:
                if referencia.startswith(prefijo):
                    referencia = referencia[len(prefijo):]
                    break
            descripcion = TIPO_MAP.get(tipo_doc.strip(), tipo_doc.strip())
            monto_texto = monto_texto.replace('.', '').replace(',', '.').replace(' ', '')
            monto = float(monto_texto.replace('-', ''))
            razon_descuento = 657 if descripcion == "Factura por convenio" else ""
            
            # ✅ Aplicar monto negativo si la razón de descuento es 657
            if razon_descuento == 657 or TIPO_MAP.get(tipo_doc.strip()) == "Nota de crédito":
                monto = -abs(monto)
            else:
                monto = -monto if '-' in monto_texto else monto

            tabla_documentos.append({
                "Tipo de documento": descripcion,
                "Referencia/ Factura": referencia,
                "Monto": monto,
                "Razon de descuento": razon_descuento
            })
    return pd.DataFrame(tabla_documentos)


def procesar(archivo_remittance, _):
    try:
        df_documentos = extraer_tabla_pdf(archivo_remittance)
        nombre_base = os.path.splitext(os.path.basename(archivo_remittance))[0]
        excel_path = os.path.join(
            os.path.dirname(archivo_remittance),
            f"Remmitance_Tottus.xlsx"
        )
        contador = 1
        while os.path.exists(excel_path):
            excel_path = os.path.join(
                os.path.dirname(archivo_remittance),
                f"Remmitance_Tottus_{contador}.xlsx"
            )
            contador += 1
        df_documentos.to_excel(excel_path, sheet_name="Documentos por pagar", index=False)
        datos_extra = [
            ("Nombre Cliente", "HIPERMERCADOS TOTTUS S A"),
            ("Numero de Cliente", "10299933"),
            ("Referencia de Pago", ""),
            ("Pago", sum(df_documentos["Monto"])),
            ("Metodo de Pago", "TRANSFERENCIA"),
            ("Fecha de pago", ""),
        ]
        wb = load_workbook(excel_path)
        ws = wb.active
        fill_azul = PatternFill(start_color="D0E9F8", end_color="D0E9F8", fill_type="solid")
        font_negrita = Font(bold=True)
        borde_negro = Border(
            left=Side(style="thin", color="000000"),
            right=Side(style="thin", color="000000"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="thin", color="000000")
        )
        for i, (col_f, col_g) in enumerate(datos_extra, start=1):
            celda_f = ws[f"F{i}"]
            celda_g = ws[f"G{i}"]
            celda_f.value = col_f
            celda_g.value = col_g
            celda_f.fill = fill_azul
            celda_f.font = font_negrita
            celda_f.border = borde_negro
            celda_g.border = borde_negro
        wb.save(excel_path)
        messagebox.showinfo("Exito", "Se genero el archivo remmitance")
    except Exception as e:
        messagebox.showerror("Error", f"Ocurrió un error: {e}")